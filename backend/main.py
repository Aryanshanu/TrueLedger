"""Cloud Run API service wrapping the agent pipeline.

Endpoints (per the build brief's API contract sketch):
    POST /borrowers/{id}/evaluate        -> triggers the pipeline, returns final decision
    GET  /borrowers/{id}/ledger          -> full ordered list of ledger steps for replay
    GET  /borrowers/{id}/consent-status  -> per-source days-remaining and confidence multiplier
    POST /borrowers/upload                -> extracts real uploaded documents into a new borrower_id
    GET  /health                         -> liveness check for Cloud Run
"""

from __future__ import annotations

import asyncio
import uuid

from fastapi import FastAPI, File, Form, HTTPException, UploadFile
from fastapi.middleware.cors import CORSMiddleware

from agents.consent import build_consent_status
from agents.extraction import assemble_fi_data, build_consent_block, extract_source
from agents.firestore_gateway import get_consent, get_decision, get_ledger, set_consent, set_fi_data
from backend.runner import evaluate_borrower

app = FastAPI(title="TrueLedger API")

# Cloud Run frontend and agent backend are separate services with separate
# URLs; wide open for the hackathon demo, tighten to the frontend's own
# origin before anything beyond a demo.
app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_methods=["*"],
    allow_headers=["*"],
)


@app.get("/health")
def health() -> dict:
    return {"status": "ok"}


@app.post("/borrowers/{borrower_id}/evaluate")
async def evaluate(borrower_id: str) -> dict:
    try:
        return await evaluate_borrower(borrower_id)
    except KeyError as exc:
        raise HTTPException(status_code=404, detail=str(exc)) from exc


@app.get("/borrowers/{borrower_id}/ledger")
def ledger(borrower_id: str) -> dict:
    entries = get_ledger(borrower_id)
    if not entries:
        raise HTTPException(status_code=404, detail=f"No ledger entries yet for {borrower_id!r} - call evaluate first.")
    return {"borrower_id": borrower_id, "steps": entries}


@app.get("/borrowers/{borrower_id}/consent-status")
def consent_status(borrower_id: str) -> dict:
    try:
        consent_data = get_consent(borrower_id)
    except KeyError as exc:
        raise HTTPException(status_code=404, detail=str(exc)) from exc

    statuses = [
        build_consent_status(source, meta["fi_type"], meta["expires_at"]).model_dump()
        for source, meta in consent_data.items()
    ]
    return {"borrower_id": borrower_id, "sources": statuses}


@app.post("/borrowers/upload")
async def upload_documents(
    loan_amount_requested: float = Form(...),
    bank_statement: UploadFile = File(...),
    gst_return: UploadFile = File(...),
    mutual_fund_statement: UploadFile | None = File(None),
    insurance_policy: UploadFile | None = File(None),
) -> dict:
    """Extracts real uploaded documents (via Gemini) into the same FI-schema
    shape the planted demo borrowers use, writes them under a brand new
    borrower_id, and returns that id - the caller then runs the existing
    POST /borrowers/{id}/evaluate exactly as it would for any other case.
    Bank statement + GST return are required (they drive the two core
    cross-checked signals); mutual fund / insurance are optional and simply
    leave that source empty, same as a real borrower who doesn't have one.
    """
    if loan_amount_requested <= 0:
        raise HTTPException(status_code=422, detail="loan_amount_requested must be greater than 0.")

    uploads: dict[str, UploadFile | None] = {
        "DEPOSIT": bank_statement,
        "GSTR1_3B": gst_return,
        "MUTUAL_FUNDS": mutual_fund_statement,
        "INSURANCE_POLICIES": insurance_policy,
    }

    # Pass 1: read every file and normalize its mime type (fast, local, no
    # network) - raises immediately on an empty file or an unsupported
    # legacy .xls, before any Gemini call is made for any document.
    to_extract: dict[str, tuple[bytes, str]] = {}
    for doc_type, upload in uploads.items():
        if upload is None:
            continue
        file_bytes = await upload.read()
        if not file_bytes or not (upload.filename or "").strip():
            if doc_type in ("MUTUAL_FUNDS", "INSURANCE_POLICIES"):
                # Not extracted_by_source[doc_type] = None here - that dict
                # doesn't exist yet (built after Pass 2 below). Simply never
                # adding this doc_type to to_extract already makes it None
                # there via its {doc_type: None for doc_type in uploads}
                # default, same as the `upload is None` branch above.
                continue
            raise HTTPException(status_code=422, detail=f"The uploaded {doc_type} file was empty.")
        mime_type = upload.content_type or "application/pdf"
        _name = (upload.filename or "").lower()
        if _name.endswith((".xlsx", ".xlsm")) or "spreadsheetml" in (mime_type or ""):
            import io
            import openpyxl
            try:
                _wb = openpyxl.load_workbook(io.BytesIO(file_bytes), read_only=True, data_only=True)
                _lines = []
                for _ws in _wb.worksheets:
                    _lines.append("# Sheet: " + str(_ws.title))
                    for _row in _ws.iter_rows(values_only=True):
                        if any(c is not None for c in _row):
                            _lines.append(",".join("" if c is None else str(c) for c in _row))
            except Exception as exc:
                raise HTTPException(
                    status_code=422, detail=f"Could not read the uploaded {doc_type} spreadsheet: {exc}"
                ) from exc
            file_bytes = "\n".join(_lines).encode("utf-8")
            mime_type = "text/plain"
        elif _name.endswith(".xls"):
            # openpyxl only reads the xlsx/xlsm zip-based format, not the
            # legacy binary .xls format - and "application/vnd.ms-excel" is
            # also the MIME type browsers legitimately send for a real .xls
            # file (not just a CSV with an Excel file association), so it
            # can't be used to route those bytes into the plain-text branch
            # below without risking sending raw binary to Gemini mislabeled
            # as text/plain.
            raise HTTPException(
                status_code=422,
                detail=f"Could not read the uploaded {doc_type} document: legacy .xls files aren't "
                "supported - please re-save as .xlsx or .csv and re-upload.",
            )
        elif _name.endswith(".csv") or mime_type in ("text/csv", "application/csv"):
            mime_type = "text/plain"
        to_extract[doc_type] = (file_bytes, mime_type)

    # Pass 2: run every document's Gemini extraction concurrently - these
    # are independent calls (one borrower's bank statement doesn't need the
    # GST return read first), so awaiting them one at a time in a loop would
    # multiply this endpoint's latency by the number of documents uploaded
    # for no reason. asyncio.gather with return_exceptions keeps one bad
    # document from cancelling extraction already in flight for the others.
    doc_types = list(to_extract.keys())
    results = await asyncio.gather(
        *(extract_source(doc_type, file_bytes, mime_type) for doc_type, (file_bytes, mime_type) in to_extract.items()),
        return_exceptions=True,
    )

    extracted_by_source: dict[str, dict | None] = {doc_type: None for doc_type in uploads}
    for doc_type, result in zip(doc_types, results):
        if isinstance(result, Exception):
            # Boundary of the system: an arbitrary user-supplied file read by
            # an LLM can fail in ways we can't enumerate (unreadable scan,
            # wrong document type, model returned malformed JSON) - turn any
            # of them into one honest 4xx rather than a 500.
            raise HTTPException(
                status_code=422, detail=f"Could not read the uploaded {doc_type} document: {result}"
            ) from result
        extracted_by_source[doc_type] = result

    # A required source that extracted to zero real data points would
    # otherwise flow silently into a pipeline that treats "no evidence" as
    # "no claim to make" (see agents/bank_agent.py's and
    # agents/gst_agent.py's own instructions: "if a metric has no evidence,
    # do not include it") rather than an error - the specialist agent would
    # report nothing, the rule table's own `if claim and other_claim`
    # guards would just never fire, and the orchestrator would still reach
    # a confident decision that never actually examined that source. Catch
    # it here, before any of that, where we can say exactly which document
    # and why - this is the one failure mode synthetic demo data can never
    # exercise, since every planted case has real non-empty data by design.
    if not extracted_by_source["DEPOSIT"]["transactions"]:
        raise HTTPException(
            status_code=422,
            detail="Could not find any transactions in the uploaded bank statement - check it's a "
            "readable statement showing individual transaction rows, not a summary page.",
        )
    if not extracted_by_source["GSTR1_3B"]["returns"]:
        raise HTTPException(
            status_code=422,
            detail="Could not find any filing periods in the uploaded GST return - check it's a "
            "GSTR-1/GSTR-3B return showing period-by-period declared turnover.",
        )

    fi_data = assemble_fi_data(extracted_by_source, loan_amount_requested)
    borrower_id = f"upload_{uuid.uuid4().hex[:10]}"
    set_fi_data(borrower_id, fi_data)
    set_consent(borrower_id, build_consent_block())
    return {"borrower_id": borrower_id}


@app.post("/ingest/onboarding")
async def ingest_onboarding(workbook: UploadFile = File(...)) -> dict:
    """Parse a 5-sheet onboarding Excel workbook and return a cross-referenced
    client graph. Every downstream sheet is joined back to Sheet 1 via
    Client_ID — invoices, GST filings, mutual fund folios, and expense records
    are all grouped under their owning client profile.

    Expected sheet layout (sheet names are matched by position, not by exact
    title, for resilience against minor naming differences):
      Sheet 1 — Client onboarding / KYB data
      Sheet 2 — Income invoices (Client_ID FK)
      Sheet 3 — GST filing periods (Client_ID FK)
      Sheet 4 — Mutual fund folios (Client_ID FK)
      Sheet 5 — Business expenses (Client_ID FK)

    Returns a summary graph plus per-client aggregates (total income, GST
    liability, portfolio value, expense total, ITC available) that the
    frontend workspace panels display without additional computation.
    """
    import io
    import openpyxl

    if not workbook.filename:
        raise HTTPException(status_code=422, detail="No file uploaded.")
    fname = workbook.filename.lower()
    if not (fname.endswith(".xlsx") or fname.endswith(".xlsm")):
        raise HTTPException(
            status_code=422,
            detail="Onboarding workbook must be an .xlsx or .xlsm file.",
        )

    raw = await workbook.read()
    if not raw:
        raise HTTPException(status_code=422, detail="Uploaded workbook is empty.")

    try:
        wb = openpyxl.load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    except Exception as exc:
        raise HTTPException(status_code=422, detail=f"Could not open workbook: {exc}") from exc

    sheets = wb.worksheets
    if len(sheets) < 1:
        raise HTTPException(status_code=422, detail="Workbook contains no sheets.")

    def _sheet_to_rows(ws) -> list[dict]:
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            return []
        headers = [str(h).strip() if h is not None else f"col_{i}" for i, h in enumerate(rows[0])]
        return [
            {headers[i]: (row[i] if i < len(row) else None) for i in range(len(headers))}
            for row in rows[1:]
            if any(v is not None for v in row)
        ]

    raw_clients = _sheet_to_rows(sheets[0]) if len(sheets) > 0 else []
    raw_invoices = _sheet_to_rows(sheets[1]) if len(sheets) > 1 else []
    raw_gst = _sheet_to_rows(sheets[2]) if len(sheets) > 2 else []
    raw_mf = _sheet_to_rows(sheets[3]) if len(sheets) > 3 else []
    raw_expenses = _sheet_to_rows(sheets[4]) if len(sheets) > 4 else []

    if not raw_clients:
        raise HTTPException(
            status_code=422,
            detail="Sheet 1 (Client Onboarding) contains no data rows. "
            "Ensure the first row is a header and subsequent rows are client records.",
        )

    # Detect the Client_ID column name (case-insensitive, underscore-tolerant)
    def _find_col(row: dict, *candidates: str) -> str | None:
        for key in row:
            normalized = key.lower().replace(" ", "_").replace("-", "_")
            for c in candidates:
                if normalized == c.lower():
                    return key
        return None

    sample = raw_clients[0] if raw_clients else {}
    cid_col = _find_col(sample, "client_id", "clientid", "client id", "id")
    if not cid_col:
        raise HTTPException(
            status_code=422,
            detail="Sheet 1 must contain a 'Client_ID' column as the primary key.",
        )

    # Build client index from Sheet 1
    client_index: dict[str, dict] = {}
    for row in raw_clients:
        cid = str(row.get(cid_col, "")).strip()
        if cid:
            client_index[cid] = {k: (str(v).strip() if v is not None else "") for k, v in row.items()}

    if not client_index:
        raise HTTPException(status_code=422, detail="No valid Client_ID values found in Sheet 1.")

    def _group_by_client(rows: list[dict]) -> dict[str, list[dict]]:
        grouped: dict[str, list[dict]] = {}
        for row in rows:
            sample_row = row
            fk_col = _find_col(sample_row, "client_id", "clientid", "client id")
            if fk_col is None:
                continue
            cid = str(row.get(fk_col, "")).strip()
            if cid:
                grouped.setdefault(cid, []).append(row)
        return grouped

    invoices_by_client = _group_by_client(raw_invoices)
    gst_by_client = _group_by_client(raw_gst)
    mf_by_client = _group_by_client(raw_mf)
    expenses_by_client = _group_by_client(raw_expenses)

    def _safe_float(v) -> float:
        try:
            return float(str(v).replace(",", "").replace("₹", "").strip()) if v is not None else 0.0
        except (ValueError, TypeError):
            return 0.0

    def _agg_invoices(rows: list[dict]) -> dict:
        total_gross = sum(_safe_float(r.get("Gross_Amount_INR") or r.get("gross_amount_inr")) for r in rows)
        total_tds = sum(_safe_float(r.get("TDS_Amount_INR") or r.get("tds_amount_inr")) for r in rows)
        total_net = sum(_safe_float(r.get("Net_Receivable_INR") or r.get("net_receivable_inr")) for r in rows)
        statuses: dict[str, int] = {}
        for r in rows:
            s = str(r.get("Status") or "").strip()
            statuses[s] = statuses.get(s, 0) + 1
        return {
            "count": len(rows),
            "total_gross_inr": round(total_gross, 2),
            "total_tds_inr": round(total_tds, 2),
            "total_net_receivable_inr": round(total_net, 2),
            "status_breakdown": statuses,
            "records": rows,
        }

    def _agg_gst(rows: list[dict]) -> dict:
        total_payable = sum(_safe_float(r.get("Net_Payable_INR") or r.get("net_payable_inr")) for r in rows)
        total_itc = sum(_safe_float(r.get("ITC_Claimed_INR") or r.get("itc_claimed_inr")) for r in rows)
        total_turnover = sum(_safe_float(r.get("Turnover_INR") or r.get("turnover_inr")) for r in rows)
        gstr1_statuses: dict[str, int] = {}
        for r in rows:
            s = str(r.get("GSTR1_Status") or r.get("gstr1_status") or "").strip()
            gstr1_statuses[s] = gstr1_statuses.get(s, 0) + 1
        return {
            "periods": len(rows),
            "total_turnover_inr": round(total_turnover, 2),
            "total_gst_payable_inr": round(total_payable, 2),
            "total_itc_claimed_inr": round(total_itc, 2),
            "net_gst_liability_inr": round(total_payable, 2),
            "gstr1_status_breakdown": gstr1_statuses,
            "records": rows,
        }

    def _agg_mf(rows: list[dict]) -> dict:
        total_value = sum(_safe_float(r.get("Current_Value_INR") or r.get("current_value_inr")) for r in rows)
        total_sip = sum(_safe_float(r.get("SIP_Amount_INR") or r.get("sip_amount_inr")) for r in rows)
        return {
            "folio_count": len(rows),
            "total_current_value_inr": round(total_value, 2),
            "total_monthly_sip_inr": round(total_sip, 2),
            "records": rows,
        }

    def _agg_expenses(rows: list[dict]) -> dict:
        total_exp = sum(_safe_float(r.get("Amount_INR") or r.get("amount_inr")) for r in rows)
        total_itc = sum(_safe_float(r.get("ITC_Amount_INR") or r.get("itc_amount_inr")) for r in rows)
        cat_totals: dict[str, float] = {}
        for r in rows:
            cat = str(r.get("Category") or "").strip() or "Other"
            cat_totals[cat] = cat_totals.get(cat, 0.0) + _safe_float(r.get("Amount_INR") or r.get("amount_inr"))
        return {
            "count": len(rows),
            "total_expense_inr": round(total_exp, 2),
            "total_itc_available_inr": round(total_itc, 2),
            "category_breakdown": {k: round(v, 2) for k, v in cat_totals.items()},
            "records": rows,
        }

    # Build the unified client graph
    clients = []
    for cid, profile in client_index.items():
        inv = _agg_invoices(invoices_by_client.get(cid, []))
        gst = _agg_gst(gst_by_client.get(cid, []))
        mf = _agg_mf(mf_by_client.get(cid, []))
        exp = _agg_expenses(expenses_by_client.get(cid, []))
        clients.append({
            "client_id": cid,
            "profile": profile,
            "summary": {
                "total_gross_income_inr": inv["total_gross_inr"],
                "total_tds_deducted_inr": inv["total_tds_inr"],
                "net_gst_liability_inr": gst["net_gst_liability_inr"],
                "total_itc_available_inr": exp["total_itc_available_inr"],
                "portfolio_value_inr": mf["total_current_value_inr"],
                "total_expenses_inr": exp["total_expense_inr"],
            },
            "income": inv,
            "gst": gst,
            "mutual_funds": mf,
            "expenses": exp,
        })

    return {
        "workbook": workbook.filename,
        "client_count": len(clients),
        "sheet_counts": {
            "clients": len(raw_clients),
            "invoices": len(raw_invoices),
            "gst_filings": len(raw_gst),
            "mutual_funds": len(raw_mf),
            "expenses": len(raw_expenses),
        },
        "clients": clients,
    }


@app.get("/borrowers/{borrower_id}/decision")
def decision(borrower_id: str) -> dict:
    """Convenience endpoint: the latest stored decision without re-running the pipeline."""
    result = get_decision(borrower_id)
    if result is None:
        raise HTTPException(status_code=404, detail=f"No decision yet for {borrower_id!r} - call evaluate first.")
    return result
