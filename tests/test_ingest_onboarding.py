"""Tests for POST /ingest/onboarding - a 200-line, live, user-reachable
endpoint (called from app.js's handleWorkbookUpload()) that previously had
zero test coverage: happy path, and the malformed-workbook cases it's
already designed to reject with a clean 422."""

from __future__ import annotations

import io

import openpyxl
import pytest
from fastapi.testclient import TestClient

from backend.main import app


def _sheet(wb, title: str, headers: list[str], rows: list[list]):
    ws = wb.create_sheet(title)
    ws.append(headers)
    for row in rows:
        ws.append(row)
    return ws


def _onboarding_workbook(
    *,
    client_rows=None,
    invoice_rows=None,
    gst_rows=None,
    mf_rows=None,
    expense_rows=None,
    client_id_header="Client_ID",
) -> bytes:
    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    if client_rows is not None:
        _sheet(wb, "Clients", [client_id_header, "Name", "City"], client_rows)
    if invoice_rows is not None:
        _sheet(wb, "Invoices", ["Client_ID", "Gross_Amount_INR", "TDS_Amount_INR", "Net_Receivable_INR", "Status"], invoice_rows)
    if gst_rows is not None:
        _sheet(wb, "GST", ["Client_ID", "Turnover_INR", "Net_Payable_INR", "ITC_Claimed_INR", "GSTR1_Status"], gst_rows)
    if mf_rows is not None:
        _sheet(wb, "MF", ["Client_ID", "Current_Value_INR", "SIP_Amount_INR"], mf_rows)
    if expense_rows is not None:
        _sheet(wb, "Expenses", ["Client_ID", "Amount_INR", "ITC_Amount_INR", "Category"], expense_rows)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _post(client: TestClient, xlsx_bytes: bytes, filename: str = "onboarding.xlsx"):
    return client.post(
        "/ingest/onboarding",
        files={"workbook": (filename, xlsx_bytes, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )


def test_happy_path_aggregates_all_five_sheets_per_client():
    xlsx = _onboarding_workbook(
        client_rows=[["CLT001", "Asha Traders", "Mumbai"], ["CLT002", "Ramesh Kumar", "Delhi"]],
        invoice_rows=[
            ["CLT001", 100000, 1000, 99000, "Settled"],
            ["CLT001", 50000, 500, 49500, "Pending"],
            ["CLT002", 20000, 200, 19800, "Settled"],
        ],
        gst_rows=[
            ["CLT001", 150000, 27000, 5000, "Filed"],
            ["CLT002", 20000, 3600, 0, "Pending"],
        ],
        mf_rows=[["CLT001", 500000, 10000]],
        expense_rows=[
            ["CLT001", 30000, 2000, "Rent"],
            ["CLT001", 10000, 500, "Utilities"],
        ],
    )
    client = TestClient(app)
    resp = _post(client, xlsx)

    assert resp.status_code == 200
    body = resp.json()
    assert body["client_count"] == 2
    assert body["sheet_counts"] == {"clients": 2, "invoices": 3, "gst_filings": 2, "mutual_funds": 1, "expenses": 2}

    by_id = {c["client_id"]: c for c in body["clients"]}
    clt001 = by_id["CLT001"]
    assert clt001["income"]["count"] == 2
    assert clt001["income"]["total_gross_inr"] == pytest.approx(150000)
    assert clt001["income"]["total_net_receivable_inr"] == pytest.approx(148500)
    assert clt001["gst"]["net_gst_liability_inr"] == pytest.approx(27000)
    assert clt001["mutual_funds"]["total_current_value_inr"] == pytest.approx(500000)
    assert clt001["expenses"]["total_expense_inr"] == pytest.approx(40000)
    assert clt001["expenses"]["category_breakdown"] == {"Rent": 30000.0, "Utilities": 10000.0}

    clt002 = by_id["CLT002"]
    assert clt002["income"]["count"] == 1
    assert clt002["gst"]["net_gst_liability_inr"] == pytest.approx(3600)
    assert clt002["mutual_funds"]["folio_count"] == 0  # no MF rows for this client


def test_client_id_column_matches_case_and_separator_variants():
    xlsx = _onboarding_workbook(
        client_rows=[["CLT001", "Asha Traders", "Mumbai"]],
        invoice_rows=[], gst_rows=[], mf_rows=[], expense_rows=[],
        client_id_header="client id",  # lowercase, space instead of underscore
    )
    client = TestClient(app)
    resp = _post(client, xlsx)

    assert resp.status_code == 200
    assert resp.json()["client_count"] == 1
    assert resp.json()["clients"][0]["client_id"] == "CLT001"


def test_non_xlsx_extension_rejected_with_422():
    client = TestClient(app)
    resp = client.post(
        "/ingest/onboarding",
        files={"workbook": ("onboarding.csv", b"Client_ID,Name\nCLT001,Asha\n", "text/csv")},
    )
    assert resp.status_code == 422
    assert "xlsx" in resp.json()["detail"].lower()


def test_empty_file_rejected_with_422():
    client = TestClient(app)
    resp = client.post(
        "/ingest/onboarding",
        files={"workbook": ("onboarding.xlsx", b"", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert resp.status_code == 422
    assert "empty" in resp.json()["detail"].lower()


def test_corrupted_workbook_bytes_rejected_with_422_not_500():
    client = TestClient(app)
    resp = client.post(
        "/ingest/onboarding",
        files={"workbook": ("onboarding.xlsx", b"this is not a real xlsx file", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert resp.status_code == 422
    assert "could not open workbook" in resp.json()["detail"].lower()


def test_sheet_one_with_no_data_rows_rejected_with_422():
    xlsx = _onboarding_workbook(client_rows=[], invoice_rows=[], gst_rows=[], mf_rows=[], expense_rows=[])
    client = TestClient(app)
    resp = _post(client, xlsx)

    assert resp.status_code == 422
    assert "no data rows" in resp.json()["detail"].lower()


def test_missing_client_id_column_rejected_with_422():
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    _sheet(wb, "Clients", ["Name", "City"], [["Asha Traders", "Mumbai"]])
    buf = io.BytesIO()
    wb.save(buf)

    client = TestClient(app)
    resp = _post(client, buf.getvalue())

    assert resp.status_code == 422
    assert "client_id" in resp.json()["detail"].lower()
