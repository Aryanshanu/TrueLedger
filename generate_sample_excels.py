import json
import os
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

output_dir = os.path.join("frontend", "public", "data", "samples")
os.makedirs(output_dir, exist_ok=True)

# Single source of truth for the 5 benchmark entities' identity and
# financials - also read by data/synthetic/generate_benchmark_js.py, which
# generates the array app.js's wizard displays. These used to be two
# hand-typed copies (one here, one in app.js) that had already drifted -
# PAN/GSTIN differed for 4 of the 5 entities between the on-screen identity
# and the identity baked into that entity's own downloadable GST return
# (its ARN reference below is derived from the PAN). `turnover` here is
# the GST-declared turnover (`gst_turnover` in the shared JSON/app.js).
with open(os.path.join("data", "benchmark_entities.json"), encoding="utf-8") as f:
    _entities_src = json.load(f)
entities = [
    {
        "id": e["id"],
        "name": e["name"],
        "sector": e["sector"],
        "turnover": e["gst_turnover"],
        "bank_inflows": e["bank_inflows"],
        "net_profit": e["net_profit"],
        "pan": e["pan"],
        "gstin": e["gstin"],
        "loan": e["loan"],
        "trend_bank": e["trend_bank"],
        "trend_gst": e["trend_gst"],
        "variance": e["variance"],
        "verdict": e["verdict"],
        "score": e["final_confidence"],
        "confidence": e["model_confidence"],
        "multiplier": e["consent_multiplier"],
        "narrative": e["narrative"],
    }
    for e in _entities_src
]

header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
border_thin = Border(
    left=Side(style='thin', color='E2E8F0'),
    right=Side(style='thin', color='E2E8F0'),
    top=Side(style='thin', color='E2E8F0'),
    bottom=Side(style='thin', color='E2E8F0')
)

for ent in entities:
    # 1. GSTR Excel
    wb_gst = openpyxl.Workbook()
    ws_gst = wb_gst.active
    ws_gst.title = "GSTR-3B & 1 Reconciliation"
    ws_gst.append(["Period", "Return Type", "Taxable Turnover (INR)", "IGST (INR)", "CGST (INR)", "SGST (INR)", "Filing Status", "ARN Reference"])
    
    monthly_gst = ent["turnover"] / 12
    months = ["Apr-25", "May-25", "Jun-25", "Jul-25", "Aug-25", "Sep-25", "Oct-25", "Nov-25", "Dec-25", "Jan-26", "Feb-26", "Mar-26"]
    for idx, m in enumerate(months):
        turnover = monthly_gst * (0.95 + (idx % 4) * 0.03)
        ws_gst.append([
            m,
            "GSTR-3B",
            round(turnover, 2),
            round(turnover * 0.10, 2),
            round(turnover * 0.04, 2),
            round(turnover * 0.04, 2),
            "Filed & Reconciled",
            f"AA{ent['pan'][:5]}{idx+10}2026"
        ])
    
    for cell in ws_gst[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")
    
    for row in ws_gst.iter_rows(min_row=2, max_col=8):
        for cell in row:
            cell.border = border_thin
            if cell.column in [3, 4, 5, 6]:
                cell.number_format = '#,##0.00'
    
    gst_path = os.path.join(output_dir, f"{ent['id']}_gstr_returns.xlsx")
    wb_gst.save(gst_path)

    # 2. ITR Excel
    wb_itr = openpyxl.Workbook()
    ws_itr = wb_itr.active
    ws_itr.title = "ITR Verification Ledger"
    ws_itr.append(["Assessment Year", "Form Type", "Gross Income (INR)", "Taxable Profit (INR)", "Depreciation (INR)", "Tax Paid (INR)", "Verification Status", "E-Filing Ack No"])
    
    ay_list = ["AY 2024-25", "AY 2025-26", "AY 2026-27"]
    for idx, ay in enumerate(ay_list):
        profit = ent["net_profit"] * (0.90 + idx * 0.12)
        gross = ent["turnover"] * (0.88 + idx * 0.10)
        ws_itr.append([
            ay,
            "ITR-6 (Corporate)",
            round(gross, 2),
            round(profit, 2),
            round(profit * 0.18, 2),
            round(profit * 0.25, 2),
            "E-Verified (CBDT API)",
            f"98421098{idx}2026"
        ])
    
    for cell in ws_itr[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")
    
    for row in ws_itr.iter_rows(min_row=2, max_col=8):
        for cell in row:
            cell.border = border_thin
            if cell.column in [3, 4, 5, 6]:
                cell.number_format = '#,##0.00'
    
    itr_path = os.path.join(output_dir, f"{ent['id']}_itr_filings.xlsx")
    wb_itr.save(itr_path)

    # 3. Bank Statement Excel
    wb_bank = openpyxl.Workbook()
    ws_bank = wb_bank.active
    ws_bank.title = "Commercial Account Statement"
    ws_bank.append(["Month", "Bank & Branch", "Total Credits (INR)", "Total Debits (INR)", "Avg Daily Balance (INR)", "Inward Mandate Txns", "Cheque Bounces", "AA Verification"])
    
    monthly_bank = ent["bank_inflows"] / 12
    for idx, m in enumerate(months):
        inflow = monthly_bank * (0.96 + (idx % 3) * 0.04)
        outflow = inflow * 0.91
        ws_bank.append([
            m,
            "HDFC Bank / Corporate Hub",
            round(inflow, 2),
            round(outflow, 2),
            round(inflow * 0.14, 2),
            round(68 + (idx * 5)),
            0 if ent["verdict"] == "APPROVE" else (1 if idx % 4 == 0 else 0),
            "Verified via Setu/Sahamati AA"
        ])
    
    for cell in ws_bank[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")
    
    for row in ws_bank.iter_rows(min_row=2, max_col=8):
        for cell in row:
            cell.border = border_thin
            if cell.column in [3, 4, 5]:
                cell.number_format = '#,##0.00'
    
    bank_path = os.path.join(output_dir, f"{ent['id']}_bank_statement.xlsx")
    wb_bank.save(bank_path)

print("Created all 15 Excel spreadsheets successfully in", output_dir)
